"""Prepare local paper artifacts from frozen evidence; never publish or overwrite."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[2]
SPEC = ROOT / 'docs/manuscript/paper_package_v1.json'


def digest(path: Path) -> str:
    """Hash a file without loading the full archive into memory."""
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def verified_spec() -> dict:
    """Reject replaced manuscript sources and evidence before creating output."""
    spec = json.loads(SPEC.read_text())
    for name, expected in spec['inputs'].items():
        if digest(ROOT / name) != expected:
            raise ValueError(f'Input hash mismatch: {name}')
    return spec


def fresh(path: Path) -> None:
    """Reserve a new destination, refusing existing directories and symlinks."""
    if path.exists() or path.is_symlink():
        raise FileExistsError(f'Choose a nonexistent destination: {path}')
    path.mkdir(parents=True, exist_ok=False)


def manifest(out: Path, spec: dict) -> None:
    """Record file identities and the explicitly uncommitted source state."""
    result = {
        'status': 'local_staging_candidate_not_published',
        'base_commit': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
        'source_state': 'working-tree snapshot; not a sealed committed-source release',
        'spec_sha256': digest(SPEC),
        'builder_sha256': digest(Path(__file__)),
        'inputs': spec['inputs'],
        'files': {str(p.relative_to(out)): {'sha256': digest(p), 'bytes': p.stat().st_size}
                  for p in sorted(out.rglob('*')) if p.is_file() and p.name != 'MANIFEST.json'},
    }
    (out / 'MANIFEST.json').write_text(json.dumps(result, indent=2)+'\n')


def inline(paragraph, text: str) -> None:
    """Render the manuscript's limited emphasis syntax as native Word runs."""
    text = ' '.join(text.split())
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))', text):
        if not token:
            continue
        value = token.strip('*`') if token.startswith(('*', '`')) else token
        for part in re.split(r'(_[ijk])\b', value):
            if not part:
                continue
            subscript = bool(re.fullmatch(r'_[ijk]', part))
            run = paragraph.add_run(part[1:] if subscript else part)
            run.font.subscript = subscript
            run.bold = token.startswith('**')
            run.italic = token.startswith(('`', '*')) and not token.startswith('**')


def document_base(label: str, double: bool = True):
    """Use restrained manuscript typography with continuous line/page numbering."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    # The bundled default template carries a colored Title border. Remove it,
    # along with theme font overrides, before applying manuscript styles.
    for node in list(doc.styles.element.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
    for node in doc.styles.element.iter(qn('w:rFonts')):
        for key in list(node.attrib):
            if key.endswith('Theme'):
                del node.attrib[key]
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.85)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(5)
    doc.styles['Normal'].paragraph_format.line_spacing = 2 if double else 1.2
    doc.styles['Normal'].paragraph_format.widow_control = True
    for name, size in [('Title', 17), ('Heading 1', 13), ('Heading 2', 11)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles['Title'].paragraph_format.line_spacing = 1.1
    h = sec.header.paragraphs[0]
    h.text = label
    h.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Page ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    if double:
        lines = OxmlElement('w:lnNumType')
        for k,v in {'countBy':'1','start':'1','restart':'continuous','distance':'180'}.items():
            lines.set(qn('w:'+k),v)
        sec._sectPr.append(lines)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.subject = label
    return doc


def add_table(doc, text: str) -> None:
    """Set explicit readable widths and repeat headers without split rows."""
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rows = [[v.strip() for v in line.strip().strip('|').split('|')]
            for line in text.splitlines() if line.strip() and not re.match(r'^\|[-: |]+\|$', line)]
    n = len(rows[0]); table = doc.add_table(rows=0, cols=n)
    table.autofit = False
    widths = {5:[.85,1.2,1.8,1.1,1.85],3:[2.3,2.25,2.25],4:[1.15,2.2,1.95,1.5]}[n]
    for col,width in zip(table.columns,widths): col.width=Inches(width)
    borders=OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        node=OxmlElement('w:'+name)
        for key,value in {'val':'single','sz':'4','color':'D9D9D9'}.items():node.set(qn('w:'+key),value)
        borders.append(node)
    table._tbl.tblPr.append(borders)
    for i,row in enumerate(rows):
        cells=table.add_row().cells
        trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=1
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.12
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
            if j>0 and n!=4:p.alignment=1
            inline(p,value)
            for r in p.runs:r.font.size=Pt(9.5);r.bold=(i==0)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def write_document(source: Path, destination: Path, spec: dict, supplement: bool = False) -> None:
    """Convert controlled paper sources to editable Word without fake metadata."""
    from docx.shared import Inches, Pt
    doc = document_base('SOD1 validation study | Paper body v2' if not supplement else 'SOD1 validation study | Supporting information', not supplement)
    for idx,block in enumerate(re.split(r'\n\s*\n',source.read_text().strip())):
        if block.startswith('# '):
            heading=block[2:].strip()
            if idx==0:doc.add_paragraph(heading,'Title')
            else:doc.add_heading(heading,1)
        elif block.startswith('## '):doc.add_heading(block[3:],2)
        elif block.startswith('|'):add_table(doc,block)
        elif supplement and (match:=re.match(r'\*\*Fig (S\d+)\.',block)):
            # One full figure and its caption per page; tall S3 has a constrained height.
            doc.add_page_break()
            key=match.group(1)
            p=doc.add_paragraph();inline(p,block)
            p.paragraph_format.keep_with_next=True
            from PIL import Image
            image=ROOT/spec['figures'][key]
            with Image.open(image) as im:w,h=im.size
            width=min(6.8,7.5*w/h)
            doc.add_picture(str(image),width=Inches(width))
        else:
            p=doc.add_paragraph();inline(p,block)
            if block.startswith('**Table '):p.paragraph_format.keep_with_next=True
            if block.startswith('**Fig ') or re.match(r'^\d+\. ',block):
                p.paragraph_format.keep_together=True
    doc.core_properties.title=doc.paragraphs[0].text
    doc.save(destination)


def documents(out: Path) -> None:
    """Build paper, supporting appendix and cover-letter body locally."""
    spec=verified_spec();fresh(out)
    write_document(ROOT/spec['manuscript'],out/'manuscript.docx',spec)
    write_document(ROOT/spec['supplement'],out/'S1_Appendix.docx',spec,True)
    doc=document_base('PLOS ONE | Cover-letter body draft',False)
    doc.add_paragraph('Cover letter','Title')
    for block in (ROOT/spec['cover_letter']).read_text().split('\n\n'):
        if block.strip():inline(doc.add_paragraph(),block)
    doc.save(out/'cover_letter.docx')
    # Lossless native-pixel supplementary copies; no resampling.
    for key,path in spec['figures'].items():
        if key.startswith('S'):
            target=out/'supporting_figures'/f'Fig{key}.png';target.parent.mkdir(exist_ok=True)
            shutil.copyfile(ROOT/path,target)
    manifest(out,spec)


def figures(out: Path) -> None:
    """Re-export four main plots with publication-size fonts from saved data."""
    import csv
    import importlib.util
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.figure import Figure
    from matplotlib.text import Text
    from src.analysis import methods_figures as mf, manuscript_figures as tf
    from src.analysis.diagnostic_audit import plot_audit
    from src.prep.build import load_config
    from src.analysis.reference_sensitivity import load_gate_attempt_1, load_scenarios
    from src.analysis.validate import evaluate_gate
    spec=verified_spec();fresh(out)
    cfg=load_config(ROOT/'config/pipeline.yaml')
    archive=(ROOT/cfg['forensic_review']['main_archive']).resolve()
    records=load_gate_attempt_1(archive,list(cfg['validation']['gate_subset']))
    refs=load_scenarios(ROOT/'data/reference_sensitivity_scenarios.csv')
    gate=evaluate_gate(records,{v:refs[v]['gate_of_record'] for v in records},cfg)
    oldsave=Figure.savefig
    def save(fig,path,**kwargs):
        original=fig.get_size_inches();width=min(7.5,float(original[0]))
        fig.set_size_inches(width,width*original[1]/original[0])
        for text in fig.findobj(match=Text):
            text.set_fontfamily('Arial');text.set_fontsize(min(11,max(8,text.get_fontsize())))
            text.set_text(text.get_text().replace('Pre-registered validation gate', 'Internally prespecified gate')
                          .replace('F64A excluded by closure rule','F64A excluded by hysteresis cap'))
        if Path(path).stem=='Fig3':
            fig.axes[2].set_ylabel('Paired movement (kcal/mol)',fontsize=8,fontfamily='Arial')
        if Path(path).stem=='Fig2':
            fig.axes[0].set_title('F64A folded r1\nSigned discrepancies cancel',fontsize=10,fontfamily='Arial')
        if Path(path).stem!='Fig3':fig.tight_layout(rect=(.02,.025,1,.94))
        oldsave(fig,path,dpi=300,facecolor='white',pil_kwargs={'compression':'tiff_lzw'})
        oldsave(fig,Path(path).with_suffix('.png'),dpi=180,facecolor='white')
    with plt.rc_context({'font.family':'Arial','font.size':9,'axes.titlesize':10,'axes.labelsize':9}):
        Figure.savefig=save
        try:
            mf.plot_gate_of_record(records,refs,gate,out/'Fig1.tif')
            rows=json.loads((ROOT/'docs/review/diagnostic_audit.json').read_text())['rows']
            with (ROOT/'docs/review/adjacent_discrepancies.csv').open() as stream:
                adjacent=list(csv.DictReader(stream))
            for row in adjacent:
                for k in ('replicate','from_window'):row[k]=int(row[k])
                row['signed_discrepancy_kcal']=float(row['signed_discrepancy_kcal'])
            plot_audit(rows,adjacent,out/'Fig2.tif')
            tspec=tf.load_spec(ROOT/'docs/manuscript_figures_v1.yaml')
            reports=tf.load_verified_reports(tspec,ROOT);data=tf.temporal_data(tspec,reports)
            tf.check_quoted_values(tspec,data)
            fig=tf.draw_temporal_figure(tspec,data,tf.lambda_vector(tspec,ROOT))
            fig.savefig(out/'Fig3.tif');plt.close(fig)
            module_spec=importlib.util.spec_from_file_location('reference_plot',ROOT/'scripts/reproduce_reference_sensitivity_independent.py')
            module=importlib.util.module_from_spec(module_spec);module_spec.loader.exec_module(module)
            module.plot(json.loads((ROOT/'docs/review/reference_reproduction.json').read_text()),ROOT/'data/reference_sensitivity_scenarios.csv',out/'Fig4.tif')
        finally:Figure.savefig=oldsave
    from PIL import Image
    for p in out.glob('*.tif'):
        with Image.open(p) as im:
            if not(789<=im.width<=2250 and im.height<=2625 and p.stat().st_size<10_000_000):
                raise ValueError(f'Figure dimensions or size out of range: {p}')
    manifest(out,spec)


def render(out: Path, renderer: Path, documents_dir: Path, figures_dir: Path,
           font_file: Path) -> None:
    """Render editable files and append viewing copies of the main figures."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import Paragraph
    from reportlab.lib.styles import ParagraphStyle
    from PIL import Image
    from pypdf import PdfReader, PdfWriter
    from xml.sax.saxutils import escape
    spec=verified_spec();fresh(out)
    qa=Path(tempfile.mkdtemp(prefix='sod1fep-paper-final-qa-'))
    for name in ('manuscript','S1_Appendix','cover_letter'):
        subprocess.run([sys.executable,str(renderer),str(documents_dir/(name+'.docx')),
                        '--output_dir',str(qa/name),'--emit_pdf'],check=True)
        shutil.copyfile(qa/name/(name+'.pdf'),out/(name+'.pdf'))
    pdfmetrics.registerFont(TTFont('PaperText',str(font_file)))
    canvas=Canvas(str(qa/'main_figures.pdf'),pagesize=(612,792))
    canvas.setTitle('Main figures for the SOD1 validation study')
    source=(ROOT/spec['manuscript']).read_text()
    for n in range(1,5):
        canvas.setFont('PaperText',11)
        canvas.drawString(40,759,f'Fig {n} | SOD1 validation study')
        path=figures_dir/f'Fig{n}.tif'
        with Image.open(path) as im:width,height=im.size
        scale=min(532/width,510/height);w,h=width*scale,height*scale
        canvas.drawImage(str(path),(612-w)/2,735-h,width=w,height=h)
        match=re.search(r'\*\*Fig '+str(n)+r'\..*?(?=\n\n)',source,re.S)
        caption=' '.join(match.group().replace('**','').split())
        paragraph=Paragraph(escape(caption),ParagraphStyle('caption',fontName='PaperText',fontSize=10,leading=13))
        _,ph=paragraph.wrap(532,180)
        if 718-h-ph<35:raise ValueError('Figure caption exceeds page')
        paragraph.drawOn(canvas,40,718-h-ph)
        canvas.showPage()
    canvas.save()
    writer=PdfWriter()
    writer.append(out/'manuscript.pdf');writer.append(qa/'main_figures.pdf')
    writer.add_metadata({'/Title':'SOD1 manuscript with main figure viewing copies'})
    writer.write(out/'manuscript_reading.pdf')
    (out/'render_record.json').write_text(json.dumps({
        'qa_directory':str(qa),'renderer_sha256':digest(renderer),
        'font_sha256':digest(font_file),
        'pages':{p.name:len(PdfReader(p).pages) for p in out.glob('*.pdf')},
        'note':'Rendering is automated; visual inspection is a separate delivery check.',
    },indent=2)+'\n')
    manifest(out,spec)
    print(f'Visual QA pages: {qa}')


def main() -> None:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('stage',choices=('documents','figures','render'))
    parser.add_argument('--output-dir',required=True,type=Path)
    parser.add_argument('--renderer',type=Path)
    parser.add_argument('--documents-dir',type=Path)
    parser.add_argument('--figures-dir',type=Path)
    parser.add_argument('--font-file',type=Path)
    args=parser.parse_args()
    if args.stage=='render':
        if not all((args.renderer,args.documents_dir,args.figures_dir,args.font_file)):
            parser.error('render needs --renderer, --documents-dir, --figures-dir and --font-file')
        render(args.output_dir,args.renderer,args.documents_dir,args.figures_dir,args.font_file)
    else:
        {'documents':documents,'figures':figures}[args.stage](args.output_dir)


if __name__=='__main__':main()
